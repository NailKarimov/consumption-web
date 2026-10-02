# -*- coding: utf-8 -*-
import io
import json
import os
import re
import shutil
import traceback
import uuid
from datetime import datetime
from contextlib import redirect_stdout, redirect_stderr

from flask import Flask, render_template, request, send_file, redirect, url_for, abort
from werkzeug.utils import secure_filename

from calculator import process_excel

# -----------------------------
# Config
# -----------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_ROOT = os.environ.get("DATA_ROOT", "/tmp/data")
os.makedirs(DATA_ROOT, exist_ok=True)

GCS_BUCKET = os.environ.get("GCS_BUCKET", "").strip()
MAX_HISTORY = int(os.environ.get("MAX_HISTORY", "200"))
# Support both BYHOURS_SHIFT_HOURS (new) and METER_SHIFT_HOURS (legacy)
BYHOURS_SHIFT_HOURS = int(
    os.environ.get("BYHOURS_SHIFT_HOURS") or os.environ.get("METER_SHIFT_HOURS") or "1"
)  # shift 1 hour earlier by default

HISTORY_KEY = "history.json"  # stored in bucket root
RUNS_PREFIX = "runs/"         # stored as runs/<job_id>/...

app = Flask(__name__)


# -----------------------------
# GCS helpers
# -----------------------------
def _get_bucket():
    if not GCS_BUCKET:
        abort(500, description="GCS_BUCKET is not configured. Set GCS_BUCKET env var and redeploy.")
    try:
        from google.cloud import storage
    except Exception:
        abort(500, description="google-cloud-storage is not installed in this deployment.")
    client = storage.Client()
    return client.bucket(GCS_BUCKET)


def _blob_exists(bucket, key: str) -> bool:
    b = bucket.blob(key)
    return b.exists()


def _download_text(bucket, key: str) -> str:
    return bucket.blob(key).download_as_text(encoding="utf-8")


def _upload_text(bucket, key: str, text: str, content_type: str = "text/plain") -> None:
    bucket.blob(key).upload_from_string(text, content_type=content_type)


def _download_bytes(bucket, key: str) -> bytes:
    return bucket.blob(key).download_as_bytes()


def _upload_file(bucket, key: str, filename: str, content_type: str = "application/octet-stream") -> None:
    bucket.blob(key).upload_from_filename(filename, content_type=content_type)


def _delete_prefix(bucket, prefix: str) -> None:
    # Delete all objects with prefix
    for b in bucket.list_blobs(prefix=prefix):
        b.delete()


def _guess_excel_mime(filename: str) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    if ext == ".xls":
        return "application/vnd.ms-excel"
    if ext in (".xlsx", ".xlsm", ".xltx", ".xltm"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return "application/octet-stream"


# -----------------------------
# History
# -----------------------------
def _history_load() -> list:
    bucket = _get_bucket()
    if not _blob_exists(bucket, HISTORY_KEY):
        return []
    try:
        raw = _download_text(bucket, HISTORY_KEY)
        data = json.loads(raw) if raw else []
        return data if isinstance(data, list) else []
    except Exception:
        return []


def _history_save(items: list) -> None:
    bucket = _get_bucket()
    # keep last MAX_HISTORY newest
    items = items[:MAX_HISTORY]
    _upload_text(bucket, HISTORY_KEY, json.dumps(items, ensure_ascii=False, indent=2), content_type="application/json")


def _parse_log_metrics(log_text: str, fallback_source_name: str = None):
    # Month
    month = "-"
    if fallback_source_name:
        m = re.search(r"(\d{4})(\d{2})(\d{2})", fallback_source_name)
        if m:
            month = "%s-%s" % (m.group(1), m.group(2))

    generated = "-"
    duration = "-"

    try:
        if log_text:
            m = re.search(r"Start(?:\s*\(local\))?\s*=\s*(\d{4}-\d{2}-\d{2})", log_text)
            if m:
                dt = datetime.strptime(m.group(1), "%Y-%m-%d")
                month = dt.strftime("%Y-%m")

            g = re.search(r"Total generated sum\s*=\s*([0-9.+\-eE]+)", log_text)
            if g:
                generated = "%.2f EUR" % float(g.group(1))

            d = re.search(r"Duration\s*=\s*([0-9]+)\s*s", log_text)
            if d:
                duration = "%d s" % int(d.group(1))
    except Exception:
        pass

    return month, generated, duration


# -----------------------------
# ByHours builder (from ByDays)
# -----------------------------
def add_byhours_from_bydays(output_xlsx_path: str) -> None:
    """
    Build ByHours sheet with the SAME structure as ByDays, but in 1-hour steps:

    Table 1 (Prices): hourly price index = mean of 4x 15-min values.
    Table 2 (Energy): sum of 4x 15-min values.
    Table 3 (Money):  sum of 4x 15-min values.

    Additionally:
    - For Energy and Money tables add RIGHT Total column (sum across dates per row).
    - Bottom-right Total×Total also filled.
    - Shift values 1 hour EARLIER with carry to NEXT DAY column:
        dest 23:00 takes src 00:00 from next date column.
    """
    try:
        from openpyxl import load_workbook
    except Exception:
        return

    wb = load_workbook(output_xlsx_path)
    if "ByDays" not in wb.sheetnames:
        return

    ws = wb["ByDays"]

    # Detect date columns by header dd.mm.yyyy in row 1
    date_cols = []
    for c in range(2, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h is None:
            continue
        hs = str(h).strip()
        if re.match(r"^\d{2}\.\d{2}\.\d{4}$", hs):
            date_cols.append((c, hs))
    if not date_cols:
        return

    def tstr(val):
        if val is None:
            return ""
        if hasattr(val, "strftime"):
            return val.strftime("%H:%M")
        s = str(val).strip()
        return s[:5]

    def find_blank_after(start_row: int) -> int:
        r = start_row
        while r <= ws.max_row:
            v = ws.cell(r, 1).value
            if v is None or str(v).strip() == "":
                return r
            r += 1
        return ws.max_row + 1

    def find_total_after(start_row: int) -> int:
        r = start_row
        while r <= ws.max_row:
            v = ws.cell(r, 1).value
            if v is not None and str(v).strip().lower() == "total":
                return r
            r += 1
        return -1

    def next_nonempty_after(start_row: int) -> int:
        r = start_row
        while r <= ws.max_row:
            v = ws.cell(r, 1).value
            if v is not None and str(v).strip() != "":
                return r
            r += 1
        return -1

    def build_time_to_row(start_row: int, end_row: int):
        m = {}
        for rr in range(start_row, end_row + 1):
            tv = tstr(ws.cell(rr, 1).value)
            if tv:
                m[tv] = rr
        return m

    def get_number(rr: int, cc: int):
        v = ws.cell(rr, cc).value
        if isinstance(v, (int, float)):
            return float(v)
        if v is None:
            return None
        try:
            return float(str(v).replace(",", "."))
        except Exception:
            return None

    # Locate 3 tables in ByDays
    prices_start = 2
    blank1 = find_blank_after(prices_start)
    prices_end = blank1 - 1
    if prices_end < prices_start:
        return

    energy_start = next_nonempty_after(blank1 + 1)
    if energy_start < 0:
        return
    energy_total_row = find_total_after(energy_start)
    if energy_total_row < 0:
        return
    energy_end = energy_total_row - 1

    blank2 = find_blank_after(energy_total_row + 1)
    money_start = next_nonempty_after(blank2 + 1)
    if money_start < 0:
        return
    money_total_row = find_total_after(money_start)
    if money_total_row < 0:
        return
    money_end = money_total_row - 1

    prices_map = build_time_to_row(prices_start, prices_end)
    energy_map = build_time_to_row(energy_start, energy_end)
    money_map = build_time_to_row(money_start, money_end)

    if "ByHours" in wb.sheetnames:
        del wb["ByHours"]
    wh = wb.create_sheet("ByHours")

    # Header row
    wh.cell(1, 1).value = "Time"
    for j, (_, hdr) in enumerate(date_cols, start=2):
        wh.cell(1, j).value = hdr

    def hour_slots(hh: str):
        return [f"{hh}:00", f"{hh}:15", f"{hh}:30", f"{hh}:45"]

    def _pick_map(kind: str):
        if kind == "prices":
            return prices_map
        if kind == "energy":
            return energy_map
        return money_map

    def write_hour_table(out_row_start: int, kind: str, add_right_total: bool, add_total_row: bool):
        """
        kind:
          'prices' -> mean of 4 values (rounded 2)
          'energy' -> sum of 4 values
          'money'  -> sum of 4 values (rounded 6)

        add_right_total: add Total column to the right (energy/money)
        add_total_row: add "Total" row at bottom (energy/money). For prices -> False (match ByDays).
        """
        map_src = _pick_map(kind)
        last_date_col = 1 + len(date_cols)
        total_col = last_date_col + 1

        if add_right_total:
            wh.cell(1, total_col).value = "Total"

        SHIFT = BYHOURS_SHIFT_HOURS

        # Fill 24 rows by DEST hour, take SRC hour/date with carry
        for dest_hour in range(24):
            # source is dest + SHIFT (shift earlier means take from later hour)
            src_total = dest_hour + SHIFT
            src_hour = src_total % 24
            date_off = src_total // 24  # 0 or 1 (for SHIFT=1 it's only 0/1)

            out_r = out_row_start + dest_hour
            wh.cell(out_r, 1).value = f"{dest_hour:02d}:00"

            slots = hour_slots(f"{src_hour:02d}")

            # write each date column, but read from shifted date column if needed
            for date_idx in range(len(date_cols)):
                out_c = 2 + date_idx
                src_date_idx = date_idx + date_off
                if src_date_idx >= len(date_cols):
                    wh.cell(out_r, out_c).value = None
                    continue

                src_col = date_cols[src_date_idx][0]  # source column in ByDays

                vals = []
                for s in slots:
                    rr = map_src.get(s)
                    if rr:
                        num = get_number(rr, src_col)
                        if num is not None:
                            vals.append(num)

                if not vals:
                    wh.cell(out_r, out_c).value = None
                else:
                    if kind == "prices":
                        wh.cell(out_r, out_c).value = round(sum(vals) / len(vals), 2)
                    elif kind == "money":
                        wh.cell(out_r, out_c).value = round(sum(vals), 6)
                    else:
                        wh.cell(out_r, out_c).value = sum(vals)

            # right total for this row (sum across dates)
            if add_right_total:
                row_vals = []
                for cc in range(2, 2 + len(date_cols)):
                    v = wh.cell(out_r, cc).value
                    if isinstance(v, (int, float)):
                        row_vals.append(float(v))
                if row_vals:
                    s = sum(row_vals)
                    wh.cell(out_r, total_col).value = round(s, 6) if kind == "money" else s
                else:
                    wh.cell(out_r, total_col).value = None

        if not add_total_row:
            # blank separator row is returned (caller can leave it empty)
            return out_row_start + 24

        # Total row (sum down the 24 rows)
        total_r = out_row_start + 24
        wh.cell(total_r, 1).value = "Total"

        for cc in range(2, 2 + len(date_cols)):
            col_vals = []
            for rr in range(out_row_start, out_row_start + 24):
                v = wh.cell(rr, cc).value
                if isinstance(v, (int, float)):
                    col_vals.append(float(v))
            if col_vals:
                s = sum(col_vals)
                wh.cell(total_r, cc).value = round(s, 6) if kind == "money" else s
            else:
                wh.cell(total_r, cc).value = None

        if add_right_total:
            grand = 0.0
            ok = False
            for rr in range(out_row_start, out_row_start + 24):
                v = wh.cell(rr, total_col).value
                if isinstance(v, (int, float)):
                    grand += float(v)
                    ok = True
            wh.cell(total_r, total_col).value = round(grand, 6) if (ok and kind == "money") else (grand if ok else None)

        return total_r + 1

    # Prices table (no total row, no right total)
    sep1 = write_hour_table(out_row_start=2, kind="prices", add_right_total=False, add_total_row=False)
    # sep1 is blank separator row (row 26 if start=2)
    # Ensure separator blank
    for cc in range(1, 2 + len(date_cols)):
        wh.cell(sep1, cc).value = None

    # Energy table (has total row and right totals)
    energy_start_out = sep1 + 1
    sep2 = write_hour_table(out_row_start=energy_start_out, kind="energy", add_right_total=True, add_total_row=True)

    # blank separator after energy
    total_col_energy = (1 + len(date_cols)) + 1
    for cc in range(1, total_col_energy + 1):
        wh.cell(sep2, cc).value = None

    # Money table (has total row and right totals)
    money_start_out = sep2 + 1
    write_hour_table(out_row_start=money_start_out, kind="money", add_right_total=True, add_total_row=True)

    if "ByDays" in wb.sheetnames:
        del wb["ByDays"]

    wb.save(output_xlsx_path)


# -----------------------------
# Routes
# -----------------------------
@app.route("/", methods=["GET"])
def index():
    jobs = _history_load()
    return render_template("index.html", jobs=jobs)


@app.route("/upload", methods=["POST"])
def upload():
    bucket = _get_bucket()

    f = request.files.get("file")
    if not f or not f.filename:
        return redirect(url_for("index"))

    job_id = datetime.utcnow().strftime("%Y%m%d-%H%M%S") + "-" + uuid.uuid4().hex[:6]
    safe_name = secure_filename(f.filename) or "upload.xlsx"

    job_dir = os.path.join(DATA_ROOT, job_id)
    input_dir = os.path.join(job_dir, "input")
    output_dir = os.path.join(job_dir, "output")
    os.makedirs(input_dir, exist_ok=True)
    os.makedirs(output_dir, exist_ok=True)

    in_path = os.path.join(input_dir, safe_name)
    out_name = "processed_" + safe_name
    out_path = os.path.join(output_dir, out_name)
    log_path = os.path.join(job_dir, "log.txt")

    f.save(in_path)

    created = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")

    # GCS keys
    run_prefix = RUNS_PREFIX + job_id + "/"
    source_key = run_prefix + "source/" + safe_name
    output_key = run_prefix + "output/" + out_name
    log_key = run_prefix + "log.txt"

    # store original in GCS
    try:
        _upload_file(bucket, source_key, in_path, content_type=_guess_excel_mime(safe_name))
    except Exception:
        source_key = None

    status = "Error"
    duration_s = 0
    log_text = ""

    t0 = datetime.utcnow()
    with open(log_path, "w", encoding="utf-8") as log:
        try:
            with redirect_stdout(log), redirect_stderr(log):
                process_excel(in_path, out_path)
            status = "Success"
        except Exception:
            traceback.print_exc(file=log)
            status = "Error"

        # Add ByHours after processing
        try:
            add_byhours_from_bydays(out_path)
        except Exception:
            traceback.print_exc(file=log)

        duration_s = int(round((datetime.utcnow() - t0).total_seconds()))
        log.write("\nDuration = %d s\n" % duration_s)

    try:
        log_text = open(log_path, "r", encoding="utf-8", errors="ignore").read()
    except Exception:
        log_text = ""

    # upload output & log
    output_file = None
    if status == "Success" and os.path.isfile(out_path):
        try:
            _upload_file(bucket, output_key, out_path, content_type=_guess_excel_mime(out_name))
            output_file = out_name
        except Exception:
            output_file = None

    try:
        _upload_file(bucket, log_key, log_path, content_type="text/plain")
    except Exception:
        pass

    month, generated, duration_text = _parse_log_metrics(log_text, fallback_source_name=safe_name)
    if duration_text == "-" and duration_s:
        duration_text = "%d s" % duration_s

    # update history
    jobs = _history_load()
    jobs.insert(
        0,
        {
            "id": job_id,
            "source": safe_name,
            "source_key": source_key,
            "month": month,
            "generated": generated,
            "status": status,
            "duration": duration_text,
            "created": created,
            "output_file": output_file,
            "output_key": output_key if output_file else None,
            "log_key": log_key,
        },
    )
    _history_save(jobs)

    return redirect(url_for("index"))


@app.route("/job/<job_id>")
def job(job_id):
    bucket = _get_bucket()
    jobs = _history_load()
    job_item = next((j for j in jobs if str(j.get("id")) == str(job_id)), None)
    if not job_item:
        abort(404)

    log_text = ""
    log_key = job_item.get("log_key")
    if log_key:
        try:
            log_text = _download_text(bucket, log_key)
        except Exception:
            log_text = ""

    files = []
    if job_item.get("output_file"):
        files = [job_item["output_file"]]

    return render_template("job.html", job_id=job_id, files=files, log=log_text)


@app.route("/download/<job_id>/<path:filename>")
def download(job_id, filename):
    bucket = _get_bucket()
    jobs = _history_load()
    job_item = next((j for j in jobs if str(j.get("id")) == str(job_id)), None)
    if not job_item:
        abort(404)

    # only allow downloading the known output
    if filename != job_item.get("output_file"):
        abort(404)

    key = job_item.get("output_key")
    if not key:
        abort(404)

    try:
        data = _download_bytes(bucket, key)
    except Exception:
        abort(404)

    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name=filename,
        mimetype=_guess_excel_mime(filename),
        max_age=0,
    )


@app.route("/download-source/<job_id>")
def download_source(job_id):
    bucket = _get_bucket()
    jobs = _history_load()
    job_item = next((j for j in jobs if str(j.get("id")) == str(job_id)), None)
    if not job_item:
        abort(404)

    source_name = job_item.get("source")
    source_key = job_item.get("source_key")
    if not source_name or not source_key:
        abort(404)

    try:
        data = _download_bytes(bucket, source_key)
    except Exception:
        abort(404)

    return send_file(
        io.BytesIO(data),
        as_attachment=True,
        download_name=source_name,
        mimetype=_guess_excel_mime(source_name),
        max_age=0,
    )


@app.route("/delete/<job_id>", methods=["POST"])
def delete(job_id):
    bucket = _get_bucket()

    # remove from history
    jobs = _history_load()
    jobs = [j for j in jobs if str(j.get("id")) != str(job_id)]
    _history_save(jobs)

    # delete files
    _delete_prefix(bucket, RUNS_PREFIX + str(job_id) + "/")

    return redirect(url_for("index"))


if __name__ == "__main__":
    # Local debug only; Cloud Run uses gunicorn
    port = int(os.environ.get("PORT", os.environ.get("APP_PORT", "8080")))
    app.run(host="0.0.0.0", port=port, debug=False)